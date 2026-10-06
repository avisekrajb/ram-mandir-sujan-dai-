# -*- coding: utf-8 -*-
"""Builds the report as a Word file (python-docx) and as an HTML file (printed to PDF by Chrome), from content.py."""
import html, sys, os
sys.path.insert(0, os.path.dirname(__file__))
import content as C
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_DIR = sys.argv[1]
BASE = sys.argv[2] if len(sys.argv) > 2 else "Security_and_Features_Report_RamMandir_ZeroInfinity_Infotech"
NAVY, BLUE, GREY = "0B1F3A", "1F6FEB", "555F6D"
HEAD_FILL, GREEN_FILL, AMBER_FILL, LINE = "EEF2F8", "E6F4EA", "E3EEFF", "D5DCE6"  # AMBER_FILL now holds the light blue used for "Set up at launch"
rgb = lambda h: RGBColor.from_string(h)

# ============================================================================================ Word
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.2)

def set_font(style, size, bold=None, color=None):
    f = style.font
    f.name = "Calibri"
    f.size = Pt(size)
    if bold is not None: f.bold = bold
    if color: f.color.rgb = rgb(color)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), "Calibri")
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]

st = doc.styles
set_font(st["Normal"], 10.5)
st["Normal"].paragraph_format.space_after = Pt(6)
st["Normal"].paragraph_format.line_spacing = 1.12
set_font(st["Heading 1"], 17, True, NAVY)
st["Heading 1"].paragraph_format.space_before = Pt(16); st["Heading 1"].paragraph_format.space_after = Pt(6)
set_font(st["Heading 2"], 13, True, BLUE)
st["Heading 2"].paragraph_format.space_before = Pt(10); st["Heading 2"].paragraph_format.space_after = Pt(6)
set_font(st["List Bullet"], 10.5)
st["List Bullet"].paragraph_format.space_after = Pt(3)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), LINE)
        b.append(e)
    tblPr.append(b)
    mar = OxmlElement("w:tblCellMar")
    for edge, w in (("top", 70), ("left", 110), ("bottom", 70), ("right", 110)):
        e = OxmlElement("w:" + edge); e.set(qn("w:w"), str(w)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblPr.append(mar)

def no_split(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    c = OxmlElement("w:cantSplit"); trPr.append(c)
    if header:
        h = OxmlElement("w:tblHeader"); trPr.append(h)

def cell_text(cell, text, bold=False, color=None, size=10, align=None):
    cell.text = ""
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = 1.08
    if align: para.alignment = align
    r = para.add_run(text); r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = rgb(color)

def make_table(head, rows, widths, status=None, boldfirst=True):
    t = doc.add_table(rows=1 + len(rows), cols=len(head))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders(t)
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]; cell_text(c, h, True, NAVY, 10); shade(c, HEAD_FILL)
    no_split(t.rows[0], header=True)
    for ri, row in enumerate(rows, start=1):
        no_split(t.rows[ri])
        for ci, val in enumerate(row):
            c = t.rows[ri].cells[ci]
            if status is not None and ci == status:
                good = val in C.GOOD
                cell_text(c, val, True, "1B6E3A" if good else "1F4FB0", 10, WD_ALIGN_PARAGRAPH.CENTER)
                shade(c, GREEN_FILL if good else AMBER_FILL)
            elif ci == 0 and boldfirst:
                cell_text(c, val, True, NAVY, 10)
            else:
                cell_text(c, val, False, None, 10)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def para(text, size=None, bold=False, color=None, align=None, after=None, style=None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text); r.bold = bold
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = rgb(color)
    if align is not None: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p

# ---- cover
for _ in range(5): para("", after=6)
para("ZEROINFINITY INFOTECH", 13, True, BLUE, after=0)
para("zeroinfinitytechnologies.com", 10, False, GREY, after=22)
para(C.TITLE, 30, True, NAVY, after=4)
para(C.SYSTEM, 17, False, NAVY, after=8)
para(C.SUBTITLE, 12, False, GREY, after=26)
cover = doc.add_table(rows=len(C.COVER), cols=2); cover.autofit = False; borders(cover)
for i, (k, v) in enumerate(C.COVER):
    cell_text(cover.rows[i].cells[0], k, True, NAVY, 10.5); shade(cover.rows[i].cells[0], HEAD_FILL)
    cell_text(cover.rows[i].cells[1], v, False, None, 10.5)
    cover.rows[i].cells[0].width = Cm(4.2); cover.rows[i].cells[1].width = Cm(12.4)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---- body
for block in C.B:
    kind, val = block[0], (block[1:] if block[0] == "image" else block[1])
    if kind == "h1pb":
        doc.add_heading(val, level=1).paragraph_format.page_break_before = True
    elif kind == "h1":
        doc.add_heading(val, level=1)
    elif kind == "h2":
        doc.add_heading(val, level=2)
    elif kind == "p":
        para(val)
    elif kind == "note":
        para(val, 9.5, False, GREY)
    elif kind == "lead":
        para(val, 10.5, True, NAVY, after=3)
    elif kind == "ul":
        for item in val:
            p = doc.add_paragraph(style="List Bullet")
            if isinstance(item, tuple):
                r = p.add_run(item[0]); r.bold = True
                p.add_run(item[1])
            else:
                p.add_run(item)
    elif kind == "table":
        make_table(val["head"], val["rows"], val["widths"], val["status"], val.get("boldfirst", True))
    elif kind == "toc":
        doc.add_heading("Contents", level=1)
        for item in val:
            para(item, 11, False, NAVY, after=5)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    elif kind == "image":
        fname, wcm, caption = val
        doc.add_picture(os.path.join(OUT_DIR, fname), width=Cm(wcm))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
        para(caption, 9, False, GREY, WD_ALIGN_PARAGRAPH.CENTER, after=10)
    elif kind == "sign":
        doc.add_heading(val["heading"], level=1)
        para(val["text"])
        s = doc.add_table(rows=len(val["rows"]) + 1, cols=4); s.autofit = False; borders(s)
        heads = ["Role", "Name", "Signature", "Date"]
        for i, h in enumerate(heads):
            cell_text(s.rows[0].cells[i], h, True, NAVY, 10); shade(s.rows[0].cells[i], HEAD_FILL)
        for ri, (role, who) in enumerate(val["rows"], start=1):
            cell_text(s.rows[ri].cells[0], role, True, NAVY, 10)
            cell_text(s.rows[ri].cells[1], who, False, None, 10)
            cell_text(s.rows[ri].cells[2], "", False, None, 10)
            cell_text(s.rows[ri].cells[3], "", False, None, 10)
            s.rows[ri].height = Cm(1.5)
        for row in s.rows:
            for i, w in enumerate((3.0, 6.2, 4.6, 2.8)):
                row.cells[i].width = Cm(w)
        doc.add_paragraph()
        para("Prepared by ZeroInfinity Infotech", 10.5, True, NAVY, after=0)
        para("zeroinfinitytechnologies.com", 10, False, GREY)

# ---- footer with page numbers
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run(C.FOOTER + "  |  Page "); r.font.size = Pt(8); r.font.color.rgb = rgb(GREY)
def field(par, instr):
    run = par.add_run(); run.font.size = Pt(8); run.font.color.rgb = rgb(GREY)
    for t, txt in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if t:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), t); run._r.append(el)
        elif txt == instr:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = instr; run._r.append(el)
        else:
            el = OxmlElement("w:t"); el.text = txt; run._r.append(el)
field(fp, " PAGE ")

# Word is strict about the order of child elements: put table / cell properties in schema order.
TBLPR = ["tblStyle","tblpPr","tblOverlap","bidiVisual","tblStyleRowBandSize","tblStyleColBandSize","tblW","jc","tblCellSpacing","tblInd","tblBorders","shd","tblLayout","tblCellMar","tblLook"]
TCPR = ["cnfStyle","tcW","gridSpan","hMerge","vMerge","tcBorders","shd","noWrap","tcMar","textDirection","tcFitText","vAlign","hideMark"]
def reorder(el, order):
    kids = list(el)
    key = lambda k: order.index(k.tag.split("}")[1]) if k.tag.split("}")[1] in order else len(order)
    for k in kids: el.remove(k)
    for k in sorted(kids, key=key): el.append(k)
for t in doc.tables:
    reorder(t._tbl.tblPr, TBLPR)
    for row in t.rows:
        for c in row.cells:
            reorder(c._tc.get_or_add_tcPr(), TCPR)

cp = doc.core_properties
cp.title = f"{C.TITLE} - {C.SYSTEM}"; cp.author = "ZeroInfinity Infotech"; cp.subject = C.SUBTITLE; cp.comments = "Confidential - prepared for the client"
doc.save(os.path.join(OUT_DIR, BASE + ".docx"))

# ============================================================================================ HTML (for the PDF)
e = html.escape
parts = []
parts.append(f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{e(C.TITLE)} - {e(C.SYSTEM)}</title>
<style>
@page {{ size: A4; margin: 20mm 22mm 22mm 22mm; }}
* {{ box-sizing: border-box; }}
body {{ font-family: Calibri, 'Segoe UI', Arial, sans-serif; font-size: 10.5pt; line-height: 1.38; color: #1c1c1c; margin: 0; }}
h1 {{ font-size: 17pt; color: #{NAVY}; margin: 22pt 0 7pt; break-after: avoid; }}
h2 {{ font-size: 13pt; color: #{BLUE}; margin: 14pt 0 6pt; break-after: avoid; }}
p {{ margin: 0 0 7pt; }}
.note {{ color: #{GREY}; font-size: 9.5pt; }}
.lead {{ color: #{NAVY}; font-weight: 700; margin-bottom: 3pt; }}
ul {{ margin: 0 0 8pt; padding-left: 18pt; }} li {{ margin-bottom: 3pt; }}
table {{ width: 100%; border-collapse: collapse; margin: 4pt 0 12pt; table-layout: fixed; }}
th, td {{ border: 0.5pt solid #{LINE}; padding: 4.5pt 6pt; vertical-align: top; font-size: 10pt; text-align: left; }}
th {{ background: #{HEAD_FILL}; color: #{NAVY}; }}
tr {{ break-inside: avoid; }}
td.k {{ font-weight: 700; color: #{NAVY}; }}
td.in {{ background: #{GREEN_FILL}; color: #1B6E3A; font-weight: 700; text-align: center; }}
td.rec {{ background: #{AMBER_FILL}; color: #1F4FB0; font-weight: 700; text-align: center; }}
.cover {{ break-after: page; padding-top: 38mm; }}
.cover .brand {{ color: #{BLUE}; font-size: 13pt; font-weight: 700; }} .cover .site {{ color: #{GREY}; font-size: 10pt; margin-bottom: 20pt; }}
.cover .title {{ color: #{NAVY}; font-size: 30pt; font-weight: 700; line-height: 1.1; margin-bottom: 6pt; }}
.cover .sys {{ color: #{NAVY}; font-size: 17pt; margin-bottom: 8pt; }} .cover .sub {{ color: #{GREY}; font-size: 12pt; margin-bottom: 26pt; }}
.sig td {{ height: 42pt; }}
figure {{ margin: 6pt 0 12pt; text-align: center; break-inside: avoid; }} figure img {{ width: 100%; }} figcaption {{ color: #{GREY}; font-size: 9pt; margin-top: 3pt; }}
.toc {{ break-after: page; }} .toc p {{ font-size: 11pt; color: #{NAVY}; margin: 0 0 7pt; }}
</style></head><body>""")
parts.append(f"""<div class="cover"><div class="brand">ZEROINFINITY INFOTECH</div><div class="site">zeroinfinitytechnologies.com</div>
<div class="title">{e(C.TITLE)}</div><div class="sys">{e(C.SYSTEM)}</div><div class="sub">{e(C.SUBTITLE)}</div>
<table><colgroup><col style="width:26%"><col></colgroup>""" + "".join(f'<tr><td class="k" style="background:#{HEAD_FILL}">{e(k)}</td><td>{e(v)}</td></tr>' for k, v in C.COVER) + "</table></div>")
for block in C.B:
    kind, val = block[0], (block[1:] if block[0] == "image" else block[1])
    if kind == "h1pb": parts.append(f'<h1 style="break-before:page;margin-top:0">{e(val)}</h1>')
    elif kind == "h1": parts.append(f"<h1>{e(val)}</h1>")
    elif kind == "h2": parts.append(f"<h2>{e(val)}</h2>")
    elif kind == "p": parts.append(f"<p>{e(val)}</p>")
    elif kind == "note": parts.append(f'<p class="note">{e(val)}</p>')
    elif kind == "lead": parts.append(f'<p class="lead">{e(val)}</p>')
    elif kind == "ul":
        lis = "".join(f"<li><b>{e(i[0])}</b>{e(i[1])}</li>" if isinstance(i, tuple) else f"<li>{e(i)}</li>" for i in val)
        parts.append(f"<ul>{lis}</ul>")
    elif kind == "table":
        w = val["widths"]; tot = sum(w)
        cols = "".join(f'<col style="width:{x / tot * 100:.1f}%">' for x in w)
        head = "".join(f"<th>{e(h)}</th>" for h in val["head"])
        body = ""
        for row in val["rows"]:
            tds = ""
            for ci, v in enumerate(row):
                if val["status"] is not None and ci == val["status"]:
                    tds += f'<td class="{"in" if v in C.GOOD else "rec"}">{e(v)}</td>'
                elif ci == 0 and val.get("boldfirst", True):
                    tds += f'<td class="k">{e(v)}</td>'
                else:
                    tds += f"<td>{e(v)}</td>"
            body += f"<tr>{tds}</tr>"
        parts.append(f"<table><colgroup>{cols}</colgroup><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>")
    elif kind == "toc":
        parts.append('<div class="toc"><h1>Contents</h1>' + "".join(f"<p>{e(i)}</p>" for i in val) + "</div>")
    elif kind == "image":
        parts.append(f'<figure><img src="{e(val[0])}" alt="{e(val[2])}"><figcaption>{e(val[2])}</figcaption></figure>')
    elif kind == "sign":
        parts.append(f"<h1>{e(val['heading'])}</h1><p>{e(val['text'])}</p>")
        rows = "".join(f'<tr><td class="k">{e(r)}</td><td>{e(w)}</td><td></td><td></td></tr>' for r, w in val["rows"])
        parts.append(f'<table class="sig"><colgroup><col style="width:18%"><col style="width:37%"><col style="width:28%"><col style="width:17%"></colgroup><thead><tr><th>Role</th><th>Name</th><th>Signature</th><th>Date</th></tr></thead><tbody>{rows}</tbody></table>')
        parts.append(f'<p class="lead" style="margin-top:14pt;margin-bottom:0">Prepared by ZeroInfinity Infotech</p><p class="note">zeroinfinitytechnologies.com</p>')
parts.append("</body></html>")
open(os.path.join(OUT_DIR, BASE + ".html"), "w", encoding="utf-8").write("".join(parts))
print("built", BASE)
